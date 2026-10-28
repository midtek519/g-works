"""Extract a local Word letter into a reviewable JSON draft; preserve prose."""
import argparse, hashlib, json, re
from pathlib import Path
from docx import Document
from render_letter import require_local

DATE = re.compile(r'^(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},?\s+\d{4}$',re.I)

def extract(source):
    source=require_local(source)
    doc=Document(source)
    if doc.tables: raise ValueError('Tables require manual migration; this importer preserves paragraph text only')
    paragraphs=[p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if not paragraphs or not DATE.match(paragraphs[0]):
        raise ValueError('No unambiguous leading date; prepare this letter manually')
    closing=next((i for i in range(len(paragraphs)-1,0,-1) if re.match(r'^(Sincerely|Yours sincerely|Respectfully|Best regards)[,\s]*$',paragraphs[i],re.I)),None)
    if closing is None or closing==len(paragraphs)-1: raise ValueError('No unambiguous closing and typed name')
    salutation=''; start=1; recipients=[]
    for i in range(1,min(closing,6)):
        if re.match(r'^(Dear\b|To whom it may concern)',paragraphs[i],re.I):
            salutation=paragraphs[i];recipients=paragraphs[1:i];start=i+1;break
    if not salutation and re.match(r'^(It is|I am|I write|I recommend|I have|This letter)',paragraphs[1],re.I) is None:
        raise ValueError('Cannot distinguish recipient block from prose; review manually')
    body=paragraphs[start:closing]
    if not body: raise ValueError('No body paragraphs found')
    return {'schema_version':1,'date':paragraphs[0],'recipient_lines':recipients,
            'subject':'','salutation':salutation,'running_label':'Historical recommendation migration proof',
            'layout':'balanced','closing':paragraphs[closing],
            'closing_lines':[line for p in paragraphs[closing+1:] for line in p.splitlines() if line.strip()],
            'paragraphs':body,
            'provenance':{'source_filename':source.name,'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
                          'body_start_paragraph':start,'closing_paragraph':closing,
                          'limitations':['Paragraph text preserved; character styling, floating images, headers, and handwritten signatures are not imported.','Historical wording is preserved, not verified or updated.']}}

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source',type=Path);parser.add_argument('--output',required=True,type=Path)
    args=parser.parse_args()
    try:
        output=require_local(args.output)
        if output.exists(): raise ValueError('Output exists; choose a new draft filename')
        data=extract(args.source)
        output.parent.mkdir(parents=True,exist_ok=True)
        output.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        print('Created local migration draft: '+str(output))
    except (ValueError,OSError) as exc: parser.exit(1,str(exc)+'\n')

if __name__=='__main__': main()
